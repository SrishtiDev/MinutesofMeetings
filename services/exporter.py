"""
services/exporter.py — Turn a Minutes object into a PDF or Excel file (as bytes).
"""

from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from models import Minutes

# Human-friendly names for the language codes stored in Minutes
_LANGUAGE_NAMES = {"en": "English", "hi": "Hindi", "mixed": "Hinglish (mixed)"}

# DejaVu covers Devanagari-free Unicode (curly quotes, accents); fall back to Helvetica
_FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")


def _register_fonts() -> tuple[str, str]:
    """Return (regular, bold) font names for the PDF."""
    try:
        pdfmetrics.registerFont(TTFont("DejaVu", str(_FONT_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(_FONT_DIR / "DejaVuSans-Bold.ttf")))
        return "DejaVu", "DejaVu-Bold"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


def to_pdf(minutes: Minutes, meeting_date: str | None = None) -> bytes:
    """Build a PDF of the minutes."""
    font, bold = _register_fonts()
    base = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=base["BodyText"], fontName=font, fontSize=10, leading=14)
    title = ParagraphStyle("title", parent=body, fontName=bold, fontSize=18, leading=22, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=body, fontName=bold, fontSize=13, leading=16, spaceBefore=14, spaceAfter=6)
    meta = ParagraphStyle("meta", parent=body, textColor="#555555")

    def p(text: str, style=body) -> Paragraph:
        return Paragraph(escape(text), style)

    language = _LANGUAGE_NAMES.get(minutes.detected_language, minutes.detected_language)
    meta_line = f"Date: {meeting_date}  |  Language: {language}" if meeting_date else f"Language: {language}"

    story = [p(minutes.title, title), p(meta_line, meta), p("Summary", h2), p(minutes.summary)]

    for heading, entries in (("Key points", minutes.key_points), ("Decisions", minutes.decisions)):
        story.append(p(heading, h2))
        if entries:
            story += [p(f"•  {e}") for e in entries]
        else:
            story.append(p("None", meta))

    story.append(p("Action items", h2))
    if minutes.action_items:
        rows = [[p("Task", ParagraphStyle("th", parent=body, fontName=bold)),
                 p("Owner", ParagraphStyle("th", parent=body, fontName=bold)),
                 p("Deadline", ParagraphStyle("th", parent=body, fontName=bold))]]
        for item in minutes.action_items:
            rows.append([p(item.task), p(item.owner or "Unassigned"), p(item.deadline or "Not specified")])
        table = Table(rows, colWidths=[95 * mm, 40 * mm, 35 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), "#EEEEEE"),
            ("GRID", (0, 0), (-1, -1), 0.5, "#BBBBBB"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(table)
    else:
        story.append(p("None", meta))

    buffer = BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, title=minutes.title,
                      leftMargin=20 * mm, rightMargin=20 * mm,
                      topMargin=20 * mm, bottomMargin=20 * mm).build(story)
    return buffer.getvalue()


def to_excel(minutes: Minutes, meeting_date: str | None = None) -> bytes:
    """Build an .xlsx workbook: a Summary sheet and an Action items sheet."""
    wb = Workbook()
    wrap = Alignment(wrap_text=True, vertical="top")
    header = Font(bold=True)

    ws = wb.active
    ws.title = "Summary"
    language = _LANGUAGE_NAMES.get(minutes.detected_language, minutes.detected_language)
    rows = [("Title", minutes.title), ("Date", meeting_date or ""), ("Language", language),
            ("Summary", minutes.summary)]
    rows += [("Key point", k) for k in minutes.key_points]
    rows += [("Decision", d) for d in minutes.decisions]
    for row in rows:
        ws.append(row)
    for cell in ws["A"]:
        cell.font = header
    for cell in ws["B"]:
        cell.alignment = wrap
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 100

    ws = wb.create_sheet("Action items")
    ws.append(["Task", "Owner", "Deadline"])
    for cell in ws[1]:
        cell.font = header
    for item in minutes.action_items:
        ws.append([item.task, item.owner or "Unassigned", item.deadline or "Not specified"])
    for col, width in zip("ABC", (70, 20, 16)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = wrap

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
